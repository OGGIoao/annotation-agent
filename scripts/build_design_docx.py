from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

SRC = Path('docs/ARCHITECTURE-DESIGN.docx.md')
OUT = Path('docs/ARCHITECTURE-DESIGN.docx')

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_code(doc, text):
    p = doc.add_paragraph()
    p.style = 'No Spacing'
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = 'Menlo'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Menlo')
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(55, 65, 81)
    return p

def add_inline_runs(p, text):
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Menlo'; r.font.size = Pt(9)
        else:
            p.add_run(part)

def add_bullets(doc, lines, numbered=False):
    for line in lines:
        p = doc.add_paragraph(style='List Number' if numbered else 'List Bullet')
        add_inline_runs(p, line)

def parse_blocks(lines):
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')
        if not line.strip():
            i += 1; continue
        if line.startswith('```'):
            lang = line[3:].strip()
            i += 1; buf=[]
            while i < len(lines) and not lines[i].startswith('```'):
                buf.append(lines[i].rstrip('\n')); i += 1
            if i < len(lines): i += 1
            yield ('code', '\n'.join(buf)); continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            yield ('heading', level, line[level:].strip()); i += 1; continue
        if line.startswith('|'):
            table=[]
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                row=lines[i].strip(); i += 1
                cells=[c.strip() for c in row.strip('|').split('|')]
                if all(set(c) <= set('-: ') for c in cells): continue
                table.append(cells)
            if table: yield ('table', table)
            continue
        if re.match(r'^[-*] ', line):
            items=[]
            while i < len(lines) and re.match(r'^[-*] ', lines[i].strip()):
                items.append(lines[i].strip()[2:]); i += 1
            yield ('bullets', items, False); continue
        if re.match(r'^\d+\. ', line):
            items=[]
            while i < len(lines) and re.match(r'^\d+\. ', lines[i].strip()):
                items.append(re.sub(r'^\d+\. ', '', lines[i].strip())); i += 1
            yield ('bullets', items, True); continue
        buf=[line]; i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#','```','|')) and not re.match(r'^[-*] |^\d+\. ', lines[i].strip()):
            buf.append(lines[i].rstrip('\n')); i += 1
        yield ('paragraph', ' '.join(buf))

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.65); sec.bottom_margin=Inches(0.65); sec.left_margin=Inches(0.75); sec.right_margin=Inches(0.75)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); styles['Normal'].font.size=Pt(10.5)
for name,size,color in [('Title',24,'17365D'),('Heading 1',17,'17365D'),('Heading 2',13,'2F5597'),('Heading 3',11,'4472C4')]:
    st=styles[name]; st.font.name='Aptos Display'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)

# Header/footer
header=sec.header.paragraphs[0]; header.text='annotation-agent  |  总设计说明'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(100,116,139)
footer=sec.footer.paragraphs[0]; footer.text='内部设计文档 · 以规则、契约和评测为质量门禁'; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.runs[0].font.size=Pt(8); footer.runs[0].font.color.rgb=RGBColor(100,116,139)

lines=SRC.read_text(encoding='utf-8').splitlines()
for block in parse_blocks(lines):
    kind=block[0]
    if kind=='heading':
        _,level,text=block
        p=doc.add_heading(text, level=min(level,3))
        if level==1: p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    elif kind=='paragraph':
        p=doc.add_paragraph(); add_inline_runs(p, block[1])
    elif kind=='code':
        add_code(doc, block[1])
    elif kind=='bullets':
        add_bullets(doc, block[1], block[2])
    elif kind=='table':
        rows=block[1]; table=doc.add_table(rows=0, cols=max(len(r) for r in rows)); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for ri,row in enumerate(rows):
            cells=table.add_row().cells
            for ci in range(len(cells)):
                set_cell_text(cells[ci], row[ci] if ci<len(row) else '', bold=(ri==0))
                if ri==0: shade(cells[ci], 'D9EAF7')
        doc.add_paragraph().paragraph_format.space_after=Pt(1)

# Save
doc.save(OUT)
print(OUT)
